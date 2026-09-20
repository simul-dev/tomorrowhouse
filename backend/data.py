"""Excel ingestion, provenance audit and reproducible initial candidate selection."""
from io import BytesIO
from pathlib import Path
from math import isfinite
import hashlib
import re
import zipfile
from openpyxl import load_workbook
from .schemas import Customer, Scenario, PRODUCTS, VOLUMES
from .distance import haversine_km

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "GLR_CASE_내일의집_수요정보.xlsx"
ALIASES = {
    "id": ["지역", "수요지역", "id", "customer_id", "region"],
    "population": ["인구", "population"],
    "lat": ["위도", "lat", "latitude"],
    "lon": ["경도", "lon", "longitude", "lng"],
    "large": ["대형", "대형가구", "large"],
    "small": ["소형", "소형가구", "small"],
    "premium": ["프리미엄_소형", "프리미엄", "프리미엄소형", "premium"],
}


class DataError(ValueError):
    def __init__(self, errors):
        self.errors = errors if isinstance(errors, list) else [str(errors)]
        super().__init__("; ".join(self.errors))


def _norm(value):
    return re.sub(r"[\s_-]", "", str(value or "")).lower()


def read_excel(source=SOURCE, sheet=None, mapping=None):
    """Read matching header within first 20 rows; reject ambiguous sheets/columns.

    mapping maps canonical names (id, lat, lon, large, small, premium,
    population) to exact source header labels. Population is optional.
    """
    raw = source if isinstance(source, bytes) else Path(source).read_bytes()
    if len(raw) > 10 * 1024 * 1024:
        raise DataError("Excel 파일은 10MB 이하만 지원합니다.")
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise DataError("압축 해제된 Excel 크기가 50MB를 초과합니다.")
        book = load_workbook(BytesIO(raw), data_only=False, read_only=True)
    except DataError:
        raise
    except Exception as exc:
        raise DataError(f"읽을 수 없는 XLSX 파일: {exc}") from exc
    try:
        if sheet and sheet not in book.sheetnames:
            raise DataError(f"시트 {sheet}가 없습니다.")
        found = []
        for ws in book:
            if sheet and ws.title != sheet:
                continue
            for rownum, row in enumerate(ws.iter_rows(min_row=1, max_row=min(20, ws.max_row or 20), values_only=True), 1):
                headers = [_norm(x) for x in row]
                cols = {}
                for field, aliases in ALIASES.items():
                    options = [_norm(mapping[field])] if mapping and field in mapping else [_norm(x) for x in aliases]
                    matches = [i for i, h in enumerate(headers) if h in options]
                    if len(matches) > 1:
                        raise DataError(f"{ws.title} {rownum}행: {field} 헤더가 중복됩니다.")
                    if matches:
                        cols[field] = matches[0]
                if all(k in cols for k in ("id", "lat", "lon", *PRODUCTS)):
                    found.append((ws, rownum, cols))
                    break
        if len(found) != 1:
            raise DataError("수요 시트를 하나로 식별할 수 없습니다. 지역·위도·경도·대형·소형·프리미엄 헤더 또는 명시적 시트/컬럼 매핑을 확인하세요.")
        ws, header_row, cols = found[0]
        customers, errors, seen, coordinates, warnings = [], [], set(), set(), []
        for rownum, row in enumerate(ws.iter_rows(min_row=header_row+1, values_only=True), header_row+1):
            if all(v is None for v in row):
                continue
            if len(customers) >= 500:
                raise DataError("수요지역은 최대 500개까지 지원합니다.")
            values = {k: row[v] if v < len(row) else None for k, v in cols.items()}
            row_errors = []
            for k, val in values.items():
                if val is None or (isinstance(val, str) and not val.strip()):
                    row_errors.append(f"{k} 누락")
                elif isinstance(val, str) and val.startswith("="):
                    row_errors.append(f"{k} 수식은 값으로 변환하여 업로드하세요")
            identity = str(values.get("id", "")).strip()
            if identity in seen:
                row_errors.append(f"중복 지역 ID {identity}")
            seen.add(identity)
            for k in ("lat", "lon", "population", *PRODUCTS):
                if k == "population" and k not in values:
                    values[k] = 0
                try:
                    if isinstance(values.get(k), bool):
                        raise ValueError()
                    number = float(values.get(k))
                    if not isfinite(number):
                        raise ValueError()
                    if k in ("population", *PRODUCTS) and (number < 0 or number != int(number)):
                        raise ValueError()
                    values[k] = number
                except (ValueError, TypeError, OverflowError):
                    row_errors.append(f"{k} 비정상 숫자")
            if row_errors:
                errors.append(f"{ws.title} {rownum}행: " + ", ".join(row_errors))
                continue
            try:
                customer = Customer(id=identity, name=re.sub(r"^(C_)?[SIK]_?", "", identity),
                                    lat=values["lat"], lon=values["lon"], population=int(values["population"]),
                                    demand={p: int(values[p]) for p in PRODUCTS})
                coord = (customer.lat, customer.lon)
                if coord in coordinates:
                    warnings.append(f"{rownum}행: 동일한 좌표의 다른 지역이 있습니다.")
                coordinates.add(coord)
                customers.append(customer)
            except ValueError as exc:
                errors.append(f"{ws.title} {rownum}행: {exc}")
        if errors:
            raise DataError(errors[:100])
        if not customers:
            raise DataError("유효한 수요행이 없습니다.")
        totals = {p: sum(getattr(c.demand, p) for c in customers) for p in PRODUCTS}
        audit = {"sheet": ws.title, "header_row": header_row, "customers": len(customers),
                 "population": sum(c.population for c in customers), "demand": totals,
                 "total_units": sum(totals.values()), "total_cbm": sum(totals[p]*VOLUMES[p] for p in PRODUCTS),
                 "sha256": hashlib.sha256(raw).hexdigest(), "warnings": warnings,
                 "column_mapping": {k: v+1 for k, v in cols.items()}}
        return customers, audit
    finally:
        book.close()


def generate_candidates(customers, count=10):
    allowed = sorted((c for c in customers if c.id.startswith(("C_K", "K_"))), key=lambda c: c.id)
    if not allowed:
        raise DataError("경기 후보를 생성할 수 없습니다. 경기 지역 ID에 C_K 또는 K_ 접두어를 사용하거나 JSON에서 후보를 명시하세요.")
    weights = [sum(getattr(c.demand, p)*VOLUMES[p] for p in PRODUCTS) for c in customers]
    distances = {a.id: [haversine_km(a.lat, a.lon, c.lat, c.lon) for c in customers] for a in allowed}
    chosen, nearest = [], [float("inf")]*len(customers)
    for _ in range(min(count, len(allowed))):
        best = min((a for a in allowed if a not in chosen),
                   key=lambda a: (sum(w*min(d, e) for w, d, e in zip(weights, nearest, distances[a.id])), a.id))
        chosen.append(best)
        nearest = [min(d, e) for d, e in zip(nearest, distances[best.id])]
    return [{"id": f"DC{i+1:02d}", "name": f"{c.name} DC", "lat": c.lat, "lon": c.lon,
             "fixed_cost": 20000000, "handling_cost": 50, "capacity_cbm": 700, "enabled": True}
            for i, c in enumerate(chosen)]


def default_scenario(source=SOURCE, sheet=None, mapping=None):
    customers, audit = read_excel(source, sheet, mapping)
    scenario = Scenario(customers=customers, facilities=generate_candidates(customers), vehicles=[
        {"id": "1t", "name": "1톤", "capacity_cbm": 6, "fixed_cost": 350000, "cost_per_km": 15000},
        {"id": "2.5t", "name": "2.5톤", "capacity_cbm": 14, "fixed_cost": 500000, "cost_per_km": 17000},
        {"id": "3.5t", "name": "3.5톤", "capacity_cbm": 17, "fixed_cost": 600000, "cost_per_km": 20000},
    ], constraints=[
        {"id": "max-dcs", "type": "max_dcs", "enabled": True, "value": 10},
        {"id": "premium-distance", "type": "premium_distance", "enabled": True, "value": 30},
        {"id": "capacity", "type": "capacity", "enabled": False, "value": None},
    ])
    return scenario, audit
