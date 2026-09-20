"""Local HTTP application; `uvicorn backend.api:app --host 127.0.0.1`."""
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import Any
import json
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from pydantic import BaseModel
from .data import DataError, default_scenario
from .schemas import Scenario

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "frontend" / "dist"
app = FastAPI(title="내일의집 Network Lab", version="1.0.0")
solve_lock = Lock()


@app.get("/api/health")
def health():
    return {"status": "ok", "distance": "Haversine", "solver": "SciPy / HiGHS"}


@app.get("/api/default")
def get_default():
    scenario, _ = default_scenario()
    return scenario


@app.get("/api/audit")
def get_audit():
    return default_scenario()[1]


@app.post("/api/validate")
def validate_scenario(scenario: Scenario):
    return scenario


@app.post("/api/upload")
def upload(file: UploadFile = File(...), sheet: str | None = Form(None), mapping: str | None = Form(None)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, "XLSX 파일만 지원합니다.")
    raw = file.file.read(10 * 1024 * 1024 + 1)
    try:
        columns = json.loads(mapping) if mapping else None
        if columns is not None and (not isinstance(columns, dict) or any(not isinstance(k,str) or not isinstance(v,str) for k,v in columns.items())):
            raise DataError("mapping은 컬럼명 문자열 객체여야 합니다.")
        scenario, audit = default_scenario(raw, sheet, columns)
        scenario.name = f"업로드 · {Path(file.filename).stem}"[:120]
        return {"scenario": scenario, "validation": audit}
    except (DataError, ValueError) as exc:
        raise HTTPException(422, getattr(exc, "errors", str(exc))) from exc


@app.post("/api/solve")
def optimize(scenario: Scenario):
    from .optimizer import solve
    if not solve_lock.acquire(blocking=False):
        raise HTTPException(409, "다른 최적화가 실행 중입니다. 완료 후 다시 실행하세요.")
    try:
        return solve(scenario)
    finally:
        solve_lock.release()


class ExportRequest(BaseModel):
    scenario: Scenario
    result: dict[str, Any]


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@app.post("/api/export")
def export_result(payload: ExportRequest):
    from .validation import validate_result
    if payload.result.get("status") not in ("optimal", "feasible_limit"):
        raise HTTPException(422, "실행 가능한 결과가 필요합니다.")
    try:
        check = validate_result(payload.scenario, payload.result)
    except (KeyError, TypeError, ValueError, IndexError) as exc:
        raise HTTPException(422, f"결과 형식이 유효하지 않습니다: {exc}") from exc
    if not check["passed"]:
        raise HTTPException(422, {"message": "입력·결과 정합성 검증 실패", "errors": check["errors"]})
    w = Workbook()
    summary = w.active
    summary.title = "Summary"
    summary.append(["scenario", safe_cell(payload.scenario.name)])
    summary.append(["status", payload.result["status"]])
    summary.append(["objective_KRW", payload.result["objective"]])
    summary.append(["mip_gap", payload.result.get("mip_gap")])
    summary.append(["distance_model", "Haversine; cost=one-way, distance KPI=round-trip"])
    costs = w.create_sheet("Costs")
    costs.append(["period", "days", "component", "daily_KRW"])
    allocations = w.create_sheet("Assignments")
    allocations.append(["period", "customer", "facility", "oneway_km", "large", "small", "premium", "unmet_large", "unmet_small", "unmet_premium"])
    trips = w.create_sheet("Trips")
    trips.append(["period", "facility", "customer", "group", "vehicle", "trips", "load_cbm", "total_capacity_cbm"])
    facilities = w.create_sheet("DCs")
    facilities.append(["period", "facility", "opened", "units", "load_cbm", "share"])
    for period in payload.result["periods"]:
        for key, value in period["costs"].items():
            costs.append([safe_cell(period["name"]), period["days"], key, value])
        for a in period["assignments"]:
            allocations.append([safe_cell(period["name"]), safe_cell(a["customer_id"]), safe_cell(a["facility_id"]), a["distance_km"],
                                *[a["quantities"][p] for p in ("large", "small", "premium")],
                                *[a["unmet"][p] for p in ("large", "small", "premium")]])
        for t in period["trips"]:
            trips.append([safe_cell(period["name"]), *[safe_cell(t[k]) for k in ("facility_id", "customer_id", "group", "vehicle_id")],
                          t["trips"], t["load_cbm"], t["capacity_cbm"]])
        for f in period["facilities"]:
            facilities.append([safe_cell(period["name"]), safe_cell(f["id"]), f["id"] in payload.result["open_facilities"], f["units"], f["load_cbm"], f["share"]])
    settings = w.create_sheet("Scenario")
    settings.append(["Scenario JSON (chunks)"])
    serialized = payload.scenario.model_dump_json(indent=2)
    for start in range(0, len(serialized), 30000):
        settings.append([serialized[start:start+30000]])
    for ws in w:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 22
    stream = BytesIO()
    w.save(stream)
    return Response(stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="tomorrowhouse-result.xlsx"'})


if DIST.is_dir():
    app.mount("/assets", StaticFiles(directory=DIST / "assets"), name="assets")


@app.get("/")
def index():
    if not (DIST / "index.html").exists():
        raise HTTPException(503, "먼저 frontend에서 npm ci 및 npm run build를 실행하세요.")
    return FileResponse(DIST / "index.html")
