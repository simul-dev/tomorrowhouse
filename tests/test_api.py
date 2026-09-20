from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.api import app
from backend.data import SOURCE

client = TestClient(app)


def test_default_and_source_upload_match_every_value():
    response = client.get("/api/default")
    assert response.status_code == 200
    default = response.json()
    uploaded = client.post("/api/upload", files={"file": ("source.xlsx", SOURCE.read_bytes())})
    assert uploaded.status_code == 200
    assert uploaded.json()["scenario"]["customers"] == default["customers"]
    assert uploaded.json()["validation"]["total_units"] == 27103
    book = load_workbook(SOURCE, data_only=True, read_only=True)
    for original, parsed in zip(list(book.active.values)[1:], default["customers"], strict=True):
        assert (parsed["id"], parsed["population"], parsed["lat"], parsed["lon"]) == original[:4]
        assert tuple(parsed["demand"][p] for p in ("large", "small", "premium")) == original[4:7]
    book.close()


def test_invalid_upload_and_invalid_scenario_return_errors():
    assert client.post("/api/upload", files={"file": ("bad.xlsx", b"not zip")}).status_code == 422
    assert client.post("/api/upload", files={"file": ("bad.txt", b"x")}).status_code == 422
    default = client.get("/api/default").json()
    default["parameters"]["time_limit"] = -1
    assert client.post("/api/solve", json=default).status_code == 422


def test_solve_then_export_excel_and_reject_inconsistent_result():
    scenario = client.get("/api/default").json()
    scenario["customers"] = scenario["customers"][:1]
    scenario["facilities"] = scenario["facilities"][:1]
    scenario["constraints"] = []
    scenario["parameters"]["mip_rel_gap"] = 0
    solved = client.post("/api/solve", json=scenario)
    assert solved.status_code == 200
    result = solved.json()
    assert result["status"] == "optimal"
    response = client.post("/api/export", json={"scenario": scenario, "result": result})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content), data_only=True)
    assert {"Summary", "Costs", "Assignments", "Trips", "DCs", "Scenario"} == set(book.sheetnames)
    assert book["Assignments"].max_row == 2
    book.close()
    result["objective"] += 1000000
    assert client.post("/api/export", json={"scenario": scenario, "result": result}).status_code == 422
